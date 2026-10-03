import importlib.util
from hashlib import sha256
import json
from pathlib import Path
import subprocess
import sys
from tempfile import TemporaryDirectory
import unittest

from openpyxl import Workbook


SCRIPT = Path(__file__).with_name("preview-historical-workbook.py")
SPEC = importlib.util.spec_from_file_location("workbook_preview", SCRIPT)
assert SPEC and SPEC.loader
preview = importlib.util.module_from_spec(SPEC)
SPEC.loader.exec_module(preview)


def create_workbook(path: Path, *, revised: bool) -> None:
    workbook = Workbook()
    pre = workbook.active
    pre.title = "Pré Embarque"
    pre["B4"] = "Necessity"
    pre["AZ4"] = "Rupture Risk"
    pre["B5"] = "2026-10-01" if not revised else "2026-10-02"
    pre["AZ5"] = "Low"
    pre["C5"] = "=1+1" if not revised else "=1+2"
    if revised:
        pre["BB4"] = "Auxiliary field"
        pre["BB5"] = "Observed"
        pre["B6"] = "New row"
        workbook.create_sheet("Extra")
    post = workbook.create_sheet("Pós Embarque")
    post["B4"] = "IP Number"
    post["B5"] = "NH-001/2026"
    workbook.save(path)


class WorkbookPreviewTests(unittest.TestCase):
    def test_detects_new_headers_rows_values_and_formula_changes(self):
        with TemporaryDirectory() as directory:
            baseline_path = Path(directory) / "baseline.xlsx"
            candidate_path = Path(directory) / "candidate.xlsx"
            create_workbook(baseline_path, revised=False)
            create_workbook(candidate_path, revised=True)
            before_hashes = [sha256(path.read_bytes()).hexdigest() for path in (baseline_path, candidate_path)]

            result = preview.compare_snapshots(
                preview.workbook_snapshot(baseline_path), preview.workbook_snapshot(candidate_path)
            )

            pre = result["sheets"]["Pré Embarque"]
            self.assertEqual(result["addedSheets"], ["Extra"])
            self.assertEqual(pre["headerChanges"], [{"column": "BB", "before": None, "after": "Auxiliary field"}])
            self.assertEqual(pre["addedRowPositions"], 1)
            self.assertEqual(pre["changedRowPositions"], 1)
            self.assertEqual(pre["changedCellsByColumn"], {"B": 1, "BB": 1})
            self.assertEqual(pre["formulaChanges"], 1)
            self.assertIn("C5", pre["formulaChangeSamples"])
            self.assertEqual([sha256(path.read_bytes()).hexdigest() for path in (baseline_path, candidate_path)], before_hashes)

    def test_rejects_unexpected_baseline_hash_without_writing_report(self):
        with TemporaryDirectory() as directory:
            baseline_path = Path(directory) / "baseline.xlsx"
            candidate_path = Path(directory) / "candidate.xlsx"
            report_path = Path(directory) / "comparison.workbook-preview.json"
            create_workbook(baseline_path, revised=False)
            create_workbook(candidate_path, revised=True)
            result = subprocess.run(
                [sys.executable, str(SCRIPT), str(baseline_path), str(candidate_path),
                 "--expected-baseline-sha256", "0" * 64, "--output", str(report_path)],
                capture_output=True, text=True, check=False,
            )
            self.assertNotEqual(result.returncode, 0)
            self.assertFalse(report_path.exists())

    def test_inventory_lists_headers_without_a_comparison_file(self):
        with TemporaryDirectory() as directory:
            path = Path(directory) / "candidate.xlsx"
            create_workbook(path, revised=True)
            result = subprocess.run(
                [sys.executable, str(SCRIPT), str(path), "--inventory"],
                capture_output=True, text=True, check=True,
            )
            report = json.loads(result.stdout)
            self.assertEqual(report["sheets"]["Pré Embarque"]["headers"]["BB"], "Auxiliary field")
            self.assertEqual(report["sheets"]["Pré Embarque"]["dataRows"], 2)


if __name__ == "__main__":
    unittest.main()
