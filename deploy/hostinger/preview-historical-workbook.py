#!/usr/bin/env python3
"""Compare two workbook snapshots without changing either workbook or a database.

Rows are compared by Excel position, not by inferred PO/IP identity. Formula
results are read from the workbook cache; formulas themselves are compared too.
"""

from __future__ import annotations

import argparse
from collections import Counter
from datetime import date, datetime, time
from hashlib import sha256
import json
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils import column_index_from_string, get_column_letter


SOURCE_SHEETS = ("Pré Embarque", "Pós Embarque")
HEADER_ROW = 4
FIRST_DATA_ROW = 5
EXTRACTED_COLUMN_RANGES = {
    "Pré Embarque": ("B", "AZ"),
    "Pós Embarque": ("B", "AS"),
}


def normalized(value: object) -> object:
    if isinstance(value, (date, datetime, time)):
        return value.isoformat()
    return value


def workbook_snapshot(path: Path) -> dict[str, object]:
    digest = sha256(path.read_bytes()).hexdigest()
    values = load_workbook(path, read_only=True, data_only=True)
    formulas = load_workbook(path, read_only=True, data_only=False)
    try:
        sheets = {}
        for name in SOURCE_SHEETS:
            if name not in values.sheetnames:
                continue
            value_sheet = values[name]
            formula_sheet = formulas[name]
            headers = {}
            rows = {}
            formula_cells = {}
            max_column = max(value_sheet.max_column, formula_sheet.max_column)
            max_row = max(value_sheet.max_row, formula_sheet.max_row)
            for row_number, row in enumerate(value_sheet.iter_rows(
                min_row=HEADER_ROW, max_row=max_row, max_col=max_column, values_only=True
            ), HEADER_ROW):
                cells = {get_column_letter(index): normalized(cell)
                         for index, cell in enumerate(row, 1) if cell is not None}
                if row_number == HEADER_ROW:
                    headers = cells
                elif cells:
                    rows[row_number] = cells
            for row in formula_sheet.iter_rows(
                min_row=HEADER_ROW, max_row=max_row, max_col=max_column
            ):
                for cell in row:
                    if cell.data_type == "f":
                        formula_cells[cell.coordinate] = cell.value
            sheets[name] = {"headers": headers, "rows": rows, "formulas": formula_cells}
        return {"path": str(path.resolve()), "sha256": digest,
                "sheetNames": values.sheetnames, "sheets": sheets}
    finally:
        values.close()
        formulas.close()


def compare_snapshots(baseline: dict, candidate: dict, sample_limit: int = 25) -> dict:
    result = {
        "comparison": "Excel row and cell positions; changes do not establish PO/IP identity",
        "baseline": {key: baseline[key] for key in ("path", "sha256")},
        "candidate": {key: candidate[key] for key in ("path", "sha256")},
        "addedSheets": sorted(set(candidate["sheetNames"]) - set(baseline["sheetNames"])),
        "removedSheets": sorted(set(baseline["sheetNames"]) - set(candidate["sheetNames"])),
        "sheets": {},
    }
    for name in SOURCE_SHEETS:
        old = baseline["sheets"].get(name)
        new = candidate["sheets"].get(name)
        if old is None or new is None:
            result["sheets"][name] = {"status": "missing in baseline" if old is None else "missing in candidate"}
            continue
        old_headers, new_headers = old["headers"], new["headers"]
        first_column, last_column = EXTRACTED_COLUMN_RANGES[name]
        first_imported_index = column_index_from_string(first_column)
        last_imported_index = column_index_from_string(last_column)
        outside_headers = [
            {"column": column, "header": header}
            for column, header in sorted(new_headers.items(), key=lambda item: column_index_from_string(item[0]))
            if not first_imported_index <= column_index_from_string(column) <= last_imported_index
        ]
        header_changes = [
            {"column": column, "before": old_headers.get(column), "after": new_headers.get(column)}
            for column in sorted(set(old_headers) | set(new_headers), key=column_index_from_string)
            if old_headers.get(column) != new_headers.get(column)
        ]
        old_rows, new_rows = old["rows"], new["rows"]
        old_positions, new_positions = set(old_rows), set(new_rows)
        changed_rows = 0
        changed_cells = 0
        by_column = Counter()
        samples = []
        for number in sorted(old_positions & new_positions):
            row_changed = False
            for column in sorted(set(old_rows[number]) | set(new_rows[number]), key=column_index_from_string):
                before = old_rows[number].get(column)
                after = new_rows[number].get(column)
                if before == after:
                    continue
                row_changed = True
                changed_cells += 1
                by_column[column] += 1
                if len(samples) < sample_limit:
                    samples.append({"cell": f"{column}{number}", "before": before, "after": after})
            changed_rows += int(row_changed)
        old_formulas, new_formulas = old["formulas"], new["formulas"]
        formula_changes = [
            cell for cell in set(old_formulas) | set(new_formulas)
            if old_formulas.get(cell) != new_formulas.get(cell)
        ]
        result["sheets"][name] = {
            "status": "compared", "headerChanges": header_changes,
            "currentExtractionRange": f"{first_column}:{last_column}",
            "candidateHeadersOutsideCurrentExtraction": outside_headers,
            "baselineRows": len(old_rows), "candidateRows": len(new_rows),
            "addedRowPositions": len(new_positions - old_positions),
            "removedRowPositions": len(old_positions - new_positions),
            "changedRowPositions": changed_rows, "changedCells": changed_cells,
            "changedCellsByColumn": dict(sorted(by_column.items(), key=lambda item: column_index_from_string(item[0]))),
            "formulaChanges": len(formula_changes),
            "formulaChangeSamples": sorted(formula_changes)[:sample_limit],
            "cellSamples": samples,
        }
    return result


def inventory(snapshot: dict) -> dict:
    sheets = {}
    for name, sheet in snapshot["sheets"].items():
        first_column, last_column = EXTRACTED_COLUMN_RANGES[name]
        first_imported_index = column_index_from_string(first_column)
        last_imported_index = column_index_from_string(last_column)
        outside_headers = [
            {"column": column, "header": header}
            for column, header in sorted(sheet["headers"].items(), key=lambda item: column_index_from_string(item[0]))
            if not first_imported_index <= column_index_from_string(column) <= last_imported_index
        ]
        sheets[name] = {
            "dataRows": len(sheet["rows"]),
            "headers": sheet["headers"],
            "formulaCells": len(sheet["formulas"]),
            "currentExtractionRange": f"{first_column}:{last_column}",
            "headersOutsideCurrentExtraction": outside_headers,
        }
    return {
        "path": snapshot["path"], "sha256": snapshot["sha256"],
        "sheetNames": snapshot["sheetNames"],
        "sheets": sheets,
    }


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("baseline", type=Path)
    parser.add_argument("candidate", type=Path, nargs="?")
    parser.add_argument("--inventory", action="store_true", help="Inspect one workbook without comparing")
    parser.add_argument("--expected-baseline-sha256")
    parser.add_argument("--output", type=Path, help="Write a local JSON report")
    parser.add_argument("--sample-limit", type=int, default=25)
    args = parser.parse_args()
    if not 0 <= args.sample_limit <= 100:
        parser.error("--sample-limit must be between 0 and 100")
    if not args.inventory and args.candidate is None:
        parser.error("candidate is required unless --inventory is used")
    if args.candidate and args.baseline.resolve() == args.candidate.resolve():
        parser.error("baseline and candidate must be different files")
    inputs = {args.baseline.resolve()}
    if args.candidate:
        inputs.add(args.candidate.resolve())
    if args.output and args.output.resolve() in inputs:
        parser.error("output cannot overwrite a workbook")
    baseline = workbook_snapshot(args.baseline)
    if args.expected_baseline_sha256 and baseline["sha256"] != args.expected_baseline_sha256.lower():
        parser.error("baseline SHA-256 differs from the expected value")
    if args.inventory:
        if args.candidate:
            parser.error("candidate cannot be used with --inventory")
        report = inventory(baseline)
    else:
        candidate = workbook_snapshot(args.candidate)
        report = compare_snapshots(baseline, candidate, args.sample_limit)
    rendered = json.dumps(report, ensure_ascii=False, indent=2, sort_keys=True, default=str)
    if args.output:
        args.output.write_text(rendered + "\n", encoding="utf-8")
        print(f"Report written to {args.output.resolve()}")
    else:
        print(rendered)


if __name__ == "__main__":
    main()
