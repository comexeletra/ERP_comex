#!/usr/bin/env python3
"""Import the approved 2026 workbook into the VPS PostgreSQL database.

Preflight is read-only. --apply requires the root-only migration DSN file and
promotes the complete workbook in one transaction. The script never reads a
database on the operator's computer.
"""

from __future__ import annotations

import argparse
from collections import Counter
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from hashlib import sha256
import json
import re
from pathlib import Path
from urllib.parse import parse_qsl, urlencode, urlsplit, urlunsplit
from uuid import UUID, uuid5

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter


APPROVED_SHA256 = "d2f025ce6dc53a15574126217cf2148fb875fbb41408f266d6a486aa5f0d7f44"
MAPPING_VERSION = "historical-2026-v1"
NAMESPACE = UUID("d708bc42-34ca-4d2a-b247-641201b6b2dc")
PRE_SHEET = "Pr\u00e9 Embarque"
POST_SHEET = "P\u00f3s Embarque"
PRE_HEADERS = {
    "B": "Necessity", "F": "Importer", "N": "PO Totvs", "T": "Product Code",
    "U": "Product Description", "W": "Qty", "X": "Unit Price",
    "Y": "Total Price", "Z": "Currency", "AB": "IP Number", "AZ": "Rupture Risk",
}
POST_HEADERS = {
    "B": "IP Number", "E": "Status", "F": "Importer", "V": "Freight Ccy.",
    "W": "Freight Cost", "AE": "Taxes Paid (R$)", "AP": "Fines R$",
    "AQ": "Storage R$", "AR": "Demurrage R$", "AS": "Qty Ctnr Dem",
}
COST_COLUMNS = (
    ("W", "FREIGHT", "V"),
    ("AE", "TAXES", None),
    ("AP", "FINES", None),
    ("AQ", "STORAGE", None),
    ("AR", "DEMURRAGE", None),
)


def stable_id(*parts: object) -> UUID:
    return uuid5(NAMESPACE, "|".join(str(part) for part in parts))


def text(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def raw_value(value: object) -> object:
    if isinstance(value, (date, datetime)):
        return value.isoformat()
    if isinstance(value, float) and not (float("-inf") < value < float("inf")):
        return str(value)
    return value


def raw_row(row: tuple[object, ...], last_column: int) -> dict[str, object]:
    return {
        get_column_letter(column): raw_value(row[column - 1] if column <= len(row) else None)
        for column in range(2, last_column + 1)
    }


def canonical_json(value: object) -> str:
    return json.dumps(value, ensure_ascii=False, sort_keys=True, separators=(",", ":"))


def decimal_value(value: object) -> Decimal | None:
    if value is None or text(value) == "":
        return None
    value_text = text(value)
    if re.fullmatch(r"[+-]?\d+,\d+", value_text):
        value_text = value_text.replace(",", ".")
    try:
        number = Decimal(value_text)
        if not number.is_finite() or len(number.as_tuple().digits) > 24:
            return None
        return number
    except InvalidOperation:
        return None


def date_value(value: object) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    return None


def currency_value(value: object) -> str | None:
    code = text(value).upper()
    return code if re.fullmatch(r"[A-Z]{3}", code) else None


def valid_ip(value: object) -> str:
    candidate = text(value)
    if not candidate or candidate.upper() in {"CANCELLED", "CANCELED"} or candidate.startswith("#"):
        return ""
    return candidate


def read_workbook(path: Path) -> dict[str, object]:
    digest = sha256(path.read_bytes()).hexdigest()
    if digest != APPROVED_SHA256:
        raise ValueError(f"Workbook SHA-256 differs from approved source: {digest}")

    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        for sheet_name, expected in ((PRE_SHEET, PRE_HEADERS), (POST_SHEET, POST_HEADERS)):
            sheet = workbook[sheet_name]
            for column, label in expected.items():
                if sheet[f"{column}4"].value != label:
                    raise ValueError(f"Unexpected layout in {sheet_name}!{column}4")

        batch_id = stable_id(digest, MAPPING_VERSION)
        source_rows = []
        po_by_key: dict[tuple[str, str], tuple] = {}
        process_by_key: dict[str, tuple] = {}
        observations = []
        links: dict[tuple[UUID, UUID], tuple] = {}
        costs = []
        issues = []
        pre_ip_keys = set()
        post_ip_keys = set()
        counts: Counter[str] = Counter()

        def issue(source_id: UUID, code: str, field: str | None, evidence: dict) -> None:
            issue_id = stable_id(batch_id, source_id, code, field or "")
            issues.append((issue_id, batch_id, source_id, "WARNING", code, field, evidence))

        def add_source(sheet_name: str, row_number: int, row: tuple, last_column: int) -> tuple[UUID, dict]:
            raw = raw_row(row, last_column)
            source_id = stable_id(batch_id, sheet_name, row_number)
            error_columns = []
            for column, value in raw.items():
                if isinstance(value, str) and value.startswith(("#REF!", "#VALUE!", "#N/A", "#DIV/0!", "#NUM!", "#NAME?", "#NULL!")):
                    error_columns.append(column)
                    issue(source_id, "EXCEL_CELL_ERROR", column, {"column": column, "value": value})
            source_rows.append((source_id, batch_id, sheet_name, row_number, raw, error_columns,
                                sha256(canonical_json(raw).encode("utf-8")).hexdigest()))
            return source_id, raw

        def add_process(ip: str, importer: str) -> UUID:
            key = ip.upper()
            if len(ip) > 80 or len(importer) > 120:
                raise ValueError("IP/importer exceeds database limit")
            prior = process_by_key.get(key)
            if prior and importer and prior[2] and prior[2] != importer:
                raise ValueError(f"Conflicting importers for IP {ip}")
            if not prior:
                process_by_key[key] = (stable_id(batch_id, "IP", key), ip, importer, key, None)
            elif importer and not prior[2]:
                process_by_key[key] = (prior[0], prior[1], importer, key, prior[4])
            return process_by_key[key][0]

        for number, row in enumerate(workbook[PRE_SHEET].iter_rows(min_row=5, max_row=6944, values_only=True), 5):
            source_id, raw = add_source(PRE_SHEET, number, row, 52)
            counts["pre_rows"] += 1
            importer = text(raw["F"])
            po = text(raw["N"])
            ip = valid_ip(raw["AB"])
            po_id = None
            process_id = None
            if po:
                if not importer or len(importer) > 120 or len(po) > 80:
                    raise ValueError(f"Invalid PO identity on pre row {number}")
                normalized = po.upper()
                po_key = (importer, normalized)
                if po_key not in po_by_key:
                    po_by_key[po_key] = (stable_id(batch_id, "PO", importer, normalized), importer, po, normalized)
                po_id = po_by_key[po_key][0]
                counts["po_rows"] += 1
            else:
                issue(source_id, "MISSING_PO", "N", {"importer": importer, "ipNumber": ip or None})
            if ip:
                process_id = add_process(ip, importer)
                pre_ip_keys.add(ip.upper())
            if po_id:
                numeric = {}
                for column in ("W", "X", "Y"):
                    numeric[column] = decimal_value(raw[column])
                    if raw[column] is not None and numeric[column] is None:
                        issue(source_id, "INVALID_NUMERIC", column, {"column": column, "value": text(raw[column])})
                currency = currency_value(raw["Z"])
                if raw["Z"] is not None and not currency:
                    issue(source_id, "INVALID_CURRENCY", "Z", {"value": text(raw["Z"])})
                necessity = date_value(row[1] if len(row) > 1 else None)
                if raw["B"] is not None and necessity is None:
                    issue(source_id, "INVALID_DATE", "B", {"value": text(raw["B"])})
                observations.append((stable_id(batch_id, "OBS", source_id), po_id, source_id, number,
                                     text(raw["T"]) or None, text(raw["U"]) or None,
                                     numeric["W"], numeric["X"], numeric["Y"], currency,
                                     necessity, text(raw["E"]) or None, text(raw["AB"]) or None, raw))
                if process_id:
                    links[(po_id, process_id)] = (po_id, process_id)

        for number, row in enumerate(workbook[POST_SHEET].iter_rows(min_row=5, max_row=194, values_only=True), 5):
            source_id, raw = add_source(POST_SHEET, number, row, 45)
            counts["post_rows"] += 1
            ip = valid_ip(raw["B"])
            if not ip:
                raise ValueError(f"Missing IP on post row {number}")
            importer = text(raw["F"])
            process_id = add_process(ip, importer)
            post_ip_keys.add(ip.upper())
            previous = process_by_key[ip.upper()]
            status = text(raw["E"]) or None
            if status and len(status) > 40:
                issue(source_id, "INVALID_STATUS", "E", {"value": status})
                status = None
            process_by_key[ip.upper()] = (previous[0], previous[1], previous[2], previous[3], status)
            for column, kind, currency_column in COST_COLUMNS:
                raw_amount = raw[column]
                if raw_amount is None:
                    continue
                amount = decimal_value(raw_amount)
                currency = currency_value(raw[currency_column]) if currency_column else "BRL"
                if amount is None or currency is None:
                    issue(source_id, "INVALID_COST", column, {"amount": text(raw_amount), "currency": text(raw.get(currency_column)) if currency_column else "BRL"})
                    continue
                costs.append((stable_id(batch_id, "COST", source_id, column), process_id, source_id,
                              column, kind, amount, currency))

        expected = {"pre_rows": 6940, "post_rows": 190, "po_rows": 6796}
        if any(counts[key] != value for key, value in expected.items()):
            raise ValueError(f"Source row counts differ: {dict(counts)}")
        if len(po_by_key) != 336 or len(pre_ip_keys) != 200 or len(post_ip_keys) != 190 or len(links) != 449:
            raise ValueError("PO/IP reconciliation differs from approved workbook")
        if post_ip_keys - pre_ip_keys:
            raise ValueError("Post-embarkation IP is absent from pre-embarkation")
        return {
            "sha256": digest, "batch_id": batch_id, "sources": source_rows,
            "orders": list(po_by_key.values()), "processes": list(process_by_key.values()),
            "observations": observations, "links": list(links.values()),
            "costs": costs, "issues": issues, "counts": dict(counts),
        }
    finally:
        workbook.close()


def migration_dsn(path: Path, database_name: str) -> str:
    if path.stat().st_mode & 0o077:
        raise ValueError("Migration DSN file must be readable only by its owner")
    lines = [line.removeprefix("MIGRATION_DATABASE_URL=") for line in path.read_text().splitlines()
             if line.startswith("MIGRATION_DATABASE_URL=")]
    if len(lines) != 1:
        raise ValueError("Expected one MIGRATION_DATABASE_URL")
    parts = urlsplit(lines[0])
    if (database_name != "erp_po_totvs_test" and not re.fullmatch(r"erp_importcheck_[a-z0-9_]+", database_name)):
        raise ValueError("Unsupported database target")
    if parts.path != "/erp_po_totvs_test" or parts.hostname not in {"127.0.0.1", "localhost"}:
        raise ValueError("Migration DSN does not target the selected VPS database")
    query = urlencode([(key, value) for key, value in parse_qsl(parts.query) if key != "uselibpqcompat"])
    return urlunsplit((parts.scheme, parts.netloc, "/" + database_name, query, parts.fragment))


def apply_import(data: dict[str, object], dsn_file: Path, database_name: str) -> bool:
    import psycopg
    from psycopg.types.json import Jsonb

    with psycopg.connect(migration_dsn(dsn_file, database_name)) as connection:
        with connection.transaction():
            with connection.cursor() as cursor:
                cursor.execute("SELECT pg_advisory_xact_lock(hashtextextended('import-erp-historical-2026', 0))")
                cursor.execute("SELECT state, source_row_count FROM migration.import_batch WHERE file_sha256 = %s AND mapping_version = %s",
                               (data["sha256"], MAPPING_VERSION))
                existing = cursor.fetchone()
                if existing:
                    if existing[0] != "PROMOTED" or existing[1] != len(data["sources"]):
                        raise ValueError("Existing import batch is incomplete or inconsistent")
                    return False
                for table in ("migration.import_batch", "migration.source_row", "procurement.purchase_order",
                              "procurement.po_line_observation", "imports.import_process", "costs.process_cost"):
                    cursor.execute(f"SELECT EXISTS (SELECT 1 FROM {table} LIMIT 1)")
                    if cursor.fetchone()[0]:
                        raise ValueError(f"Refusing first import: {table} is not empty")

                cursor.execute("INSERT INTO migration.import_batch (id, file_name, file_sha256, mapping_version, state, source_row_count) VALUES (%s, %s, %s, %s, 'STAGED', %s)",
                               (data["batch_id"], "Follow Up Import 2026.xlsx", data["sha256"], MAPPING_VERSION, len(data["sources"])))
                cursor.executemany("INSERT INTO migration.source_row (id, batch_id, sheet_name, row_number, raw_values, error_columns, row_hash) VALUES (%s, %s, %s, %s, %s, %s, %s)",
                                   [(a,b,c,d,Jsonb(e),Jsonb(f),g) for a,b,c,d,e,f,g in data["sources"]])
                cursor.executemany("INSERT INTO procurement.purchase_order (id, importer, external_number, normalized_number) VALUES (%s, %s, %s, %s)", data["orders"])
                cursor.executemany("INSERT INTO imports.import_process (id, ip_number, importer, normalized_ip_number, logistics_status) VALUES (%s, %s, %s, %s, %s)", data["processes"])
                cursor.executemany("INSERT INTO procurement.po_line_observation (id, purchase_order_id, source_row_id, source_row_number, product_code_snapshot, description_snapshot, quantity, unit_price, historical_amount, currency_code, necessity_date, historical_status, source_ip_text, raw_values) VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s)",
                                   [(*row[:-1], Jsonb(row[-1])) for row in data["observations"]])
                cursor.executemany("INSERT INTO procurement.process_purchase_order (purchase_order_id, process_id) VALUES (%s, %s)", data["links"])
                cursor.executemany("INSERT INTO costs.process_cost (id, process_id, source_row_id, source_column, cost_type, amount, currency_code) VALUES (%s, %s, %s, %s, %s, %s, %s)", data["costs"])
                cursor.executemany("INSERT INTO migration.data_issue (id, batch_id, source_row_id, severity, issue_code, field_name, evidence) VALUES (%s, %s, %s, %s, %s, %s, %s)",
                                   [(*row[:-1], Jsonb(row[-1])) for row in data["issues"]])
                cursor.execute("UPDATE migration.import_batch SET state = 'PROMOTED', promoted_at = now() WHERE id = %s", (data["batch_id"],))
                checks = {
                    "migration.source_row": len(data["sources"]),
                    "procurement.purchase_order": len(data["orders"]),
                    "procurement.po_line_observation": len(data["observations"]),
                    "imports.import_process": len(data["processes"]),
                    "procurement.process_purchase_order": len(data["links"]),
                    "costs.process_cost": len(data["costs"]),
                    "migration.data_issue": len(data["issues"]),
                }
                for table, expected in checks.items():
                    cursor.execute(f"SELECT count(*) FROM {table}")
                    if cursor.fetchone()[0] != expected:
                        raise ValueError(f"Reconciliation failed for {table}")
                return True


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--workbook", type=Path, required=True)
    parser.add_argument("--apply", action="store_true", help="Promote into the selected VPS database")
    parser.add_argument("--dsn-file", type=Path, default=Path("/etc/import-erp/migration-release.env"))
    parser.add_argument("--database-name", default="erp_po_totvs_test",
                        help="Selected database or an erp_importcheck_* restored copy")
    args = parser.parse_args()
    data = read_workbook(args.workbook)
    summary = {"sha256": data["sha256"], **data["counts"],
               "source_rows": len(data["sources"]), "purchase_orders": len(data["orders"]),
               "observations": len(data["observations"]), "import_processes": len(data["processes"]),
               "po_ip_links": len(data["links"]), "process_costs": len(data["costs"]),
               "quality_issues": len(data["issues"])}
    print(json.dumps(summary, ensure_ascii=False, sort_keys=True))
    if args.apply:
        print("PROMOTED" if apply_import(data, args.dsn_file, args.database_name) else "ALREADY_PROMOTED")


if __name__ == "__main__":
    main()
