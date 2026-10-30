#!/usr/bin/env python3
"""Turn the Status PO's sheet of the MRP workbook into one SQL transaction for the HML database.

Rules and decisions are in docs/PLANO_CARGA_HISTORICA_PLANILHA.md. The Follow Up decides POs and SCs;
this sheet only gives the date the commercial plan was received (column J) of each SC (column Q). The
date lives on the SC and is copied to its product lines, which is what the End to End reads.

The script never connects to a database. The SQL refuses any database other than erp_po_totvs_hml and,
without --commit, ends in ROLLBACK with every trigger active.
"""

from __future__ import annotations

import argparse
from datetime import date, datetime
from hashlib import sha256
import json
from pathlib import Path
from uuid import UUID, uuid5

NAMESPACE = UUID("8d0f7a7c-4f4b-4c55-9d0e-2b6f1c3e7a12")
ACTOR = "migration:commercial-plan-load"
EXPECTED_DATABASE = "erp_po_totvs_hml"
SHEET = "Status PO's"
HEADER_ROW = 2
PO_COLUMN, DATE_COLUMN, SC_COLUMN = 3, 9, 16
DIGITS = 6
# The Follow Up leaves this PO without SC; the user chose to create it and link its lines.
NEW_REQUESTS = {"020328": "031362"}


def text(value: object) -> str:
    return "" if value is None else str(value).strip()


def sql_text(value: object) -> str:
    return "'" + str(value).replace("'", "''") + "'"


def number(value: object, label: str, row: int) -> str:
    if isinstance(value, bool) or not isinstance(value, (int, str)) or not text(value).isdigit():
        raise ValueError(f"Row {row}: unexpected {label}: {value!r}")
    return text(value).zfill(DIGITS)


def check_headers(header: tuple) -> None:
    found = (text(header[PO_COLUMN]).upper(), " ".join(text(header[DATE_COLUMN]).upper().split()), text(header[SC_COLUMN]).upper())
    if found != ("PO", "RECEBIMENTO PLANO", "SC"):
        raise ValueError(f"Unexpected columns D, J and Q: {found}")


def build_plan(rows: list[tuple[int, tuple]]) -> list[dict]:
    """One entry per PO and SC. A row without its PO, SC or date stops the load."""
    plan: dict[tuple[str, str], str] = {}
    for row_number, cells in rows:
        received = cells[DATE_COLUMN]
        if isinstance(received, datetime):
            received = received.date()
        if not isinstance(received, date):
            raise ValueError(f"Row {row_number}: unexpected date: {received!r}")
        key = (number(cells[PO_COLUMN], "PO", row_number), number(cells[SC_COLUMN], "SC", row_number))
        value = received.isoformat()
        if plan.setdefault(key, value) != value:
            raise ValueError(f"Row {row_number}: PO {key[0]} and SC {key[1]} have two dates")
    return [{"po": po, "sc": sc, "received": received} for (po, sc), received in sorted(plan.items())]


def summary(plan: list[dict], digest: str, expected_items: int) -> dict:
    return {"sha256": digest, "actor": ACTOR, "requests": len({entry["sc"] for entry in plan}), "items": expected_items,
            "createdRequests": len(NEW_REQUESTS), "linkedOrders": len(NEW_REQUESTS), "plan": plan}


def render_new_request(sc: str, po: str, batch_id: str) -> str:
    request_id = str(uuid5(NAMESPACE, f"SC|{sc}"))
    lines = ("FROM procurement.purchase_order_item AS item WHERE item.purchase_order_id = po.id")
    return f"""DO $new_request$
DECLARE orders bigint;
BEGIN
  SELECT count(*) INTO orders FROM procurement.purchase_order AS po
   WHERE po.normalized_number = {sql_text(po)} AND po.lifecycle_status = 'ACTIVE' AND po.purchase_request_id IS NULL
     AND NOT EXISTS (SELECT 1 {lines} AND item.sc_number IS NOT NULL);
  IF orders <> 1 OR EXISTS (SELECT 1 FROM procurement.purchase_request WHERE normalized_sc_number = {sql_text(sc)}) THEN
    RAISE EXCEPTION 'PO {po} is not the active PO without SC that SC {sc} expects';
  END IF;
END $new_request$;
INSERT INTO procurement.purchase_request (id, importer, sc_number, normalized_sc_number, requester, approval_date, supplier_text, created_by)
SELECT {sql_text(request_id)}, po.importer, {sql_text(sc)}, {sql_text(sc)},
       (SELECT mode() WITHIN GROUP (ORDER BY item.requester) {lines}),
       (SELECT mode() WITHIN GROUP (ORDER BY item.sc_approval_date) {lines}),
       po.supplier_text, {sql_text(ACTOR)}
  FROM procurement.purchase_order AS po WHERE po.normalized_number = {sql_text(po)};
UPDATE procurement.purchase_order_item AS item SET sc_number = {sql_text(sc)}, updated_at = now()
  FROM procurement.purchase_order AS po WHERE item.purchase_order_id = po.id AND po.normalized_number = {sql_text(po)};
UPDATE procurement.purchase_order SET purchase_request_id = {sql_text(request_id)}, version = version + 1, updated_at = now()
 WHERE normalized_number = {sql_text(po)};
INSERT INTO audit.audit_log (id, aggregate_type, aggregate_id, entity_type, entity_id, operation, new_value, actor_id, occurred_at, reason, correlation_id)
SELECT gen_random_uuid(), 'PURCHASE_ORDER', id, 'PURCHASE_ORDER', id, 'HISTORICAL_IMPORT',
       jsonb_build_object('scNumber', {sql_text(sc)}), {sql_text(ACTOR)}, now(),
       'SC da planilha MRP vinculada à PO sem SC no Follow Up.', {sql_text(batch_id)}::uuid
  FROM procurement.purchase_order WHERE normalized_number = {sql_text(po)};"""


def render_sql(plan: list[dict], digest: str, expected_items: int, commit: bool) -> str:
    batch_id = str(uuid5(NAMESPACE, digest))
    values = ",\n".join(f"({sql_text(entry['po'])},{sql_text(entry['sc'])},{sql_text(entry['received'])}::date)" for entry in plan)
    requests = len({entry["sc"] for entry in plan})
    parts = ["\\set ON_ERROR_STOP on", "BEGIN;", "SET LOCAL lock_timeout = '30s';", f"SET LOCAL app.actor_id = {sql_text(ACTOR)};",
             f"""DO $guard$
BEGIN
  IF current_database() <> '{EXPECTED_DATABASE}' THEN
    RAISE EXCEPTION 'Refusing to load into %: only {EXPECTED_DATABASE} is allowed', current_database();
  END IF;
  IF EXISTS (SELECT 1 FROM audit.audit_log WHERE correlation_id = {sql_text(batch_id)}::uuid) THEN
    RAISE EXCEPTION 'ALREADY_PROMOTED: this workbook was already loaded';
  END IF;
  IF EXISTS (SELECT 1 FROM procurement.purchase_request WHERE commercial_plan_received_date IS NOT NULL)
     OR EXISTS (SELECT 1 FROM procurement.purchase_order_item WHERE commercial_plan_received_date IS NOT NULL) THEN
    RAISE EXCEPTION 'The load does not replace dates already recorded';
  END IF;
END $guard$;""",
             f"""CREATE TEMP TABLE plan_row (po text NOT NULL, sc text NOT NULL, received date NOT NULL, PRIMARY KEY (po, sc)) ON COMMIT DROP;
INSERT INTO plan_row (po, sc, received) VALUES
{values};"""]
    parts += [render_new_request(sc, po, batch_id) for sc, po in sorted(NEW_REQUESTS.items())]
    parts.append("""DO $known$
BEGIN
  IF EXISTS (SELECT 1 FROM plan_row AS plan WHERE (SELECT count(*) FROM procurement.purchase_request AS request
              WHERE request.normalized_sc_number = plan.sc) > 1) THEN
    RAISE EXCEPTION 'An SC of the sheet exists under more than one importer';
  END IF;
END $known$;
-- One date for the line. Two dates on its SCs: the date of its PO in the sheet, or the load stops.
CREATE TEMP TABLE plan_item ON COMMIT DROP AS
WITH item_sc AS (
  SELECT item.id, item.purchase_order_id, po.normalized_number AS po_number,
         unnest(string_to_array(replace(item.sc_number, ' ', ''), ',')) AS sc
    FROM procurement.purchase_order_item AS item
    JOIN procurement.purchase_order AS po ON po.id = item.purchase_order_id
   WHERE item.sc_number IS NOT NULL
), dated AS (
  SELECT item_sc.id, item_sc.purchase_order_id, item_sc.po_number, array_agg(DISTINCT plan.received) AS dates,
         array_agg(DISTINCT plan.received) FILTER (WHERE plan.po = item_sc.po_number) AS po_dates
    FROM item_sc JOIN plan_row AS plan ON plan.sc = item_sc.sc
   GROUP BY 1, 2, 3
)
SELECT id, purchase_order_id, po_number,
       CASE WHEN cardinality(dates) = 1 THEN dates[1] WHEN cardinality(po_dates) = 1 THEN po_dates[1] END AS received
  FROM dated;
DO $resolved$
BEGIN
  IF EXISTS (SELECT 1 FROM plan_item WHERE received IS NULL) THEN
    RAISE EXCEPTION 'A line has two dates and no date of its own PO';
  END IF;
END $resolved$;
-- The SC holds one date: with two, the one of the PO with the most lines.
CREATE TEMP TABLE plan_request ON COMMIT DROP AS
SELECT DISTINCT ON (plan.sc) request.id, plan.sc, plan.received
  FROM plan_row AS plan
  JOIN procurement.purchase_request AS request ON request.normalized_sc_number = plan.sc
 ORDER BY plan.sc, (SELECT count(*) FROM plan_item AS line WHERE line.po_number = plan.po AND line.received = plan.received) DESC, plan.received;
-- Closed and cancelled POs refuse item changes: they reopen here and return exactly as they were.
CREATE TEMP TABLE plan_po_state ON COMMIT DROP AS
SELECT po.id, po.lifecycle_status, po.lifecycle_changed_at, po.lifecycle_changed_by, po.lifecycle_reason, po.updated_at
  FROM procurement.purchase_order AS po
 WHERE po.lifecycle_status <> 'ACTIVE' AND po.id IN (SELECT purchase_order_id FROM plan_item);
UPDATE procurement.purchase_order AS po SET lifecycle_status = 'ACTIVE' FROM plan_po_state AS state WHERE po.id = state.id;
UPDATE procurement.purchase_order_item AS item SET commercial_plan_received_date = plan.received, updated_at = now()
  FROM plan_item AS plan WHERE item.id = plan.id;
UPDATE procurement.purchase_order AS po
   SET lifecycle_status = state.lifecycle_status, lifecycle_changed_at = state.lifecycle_changed_at,
       lifecycle_changed_by = state.lifecycle_changed_by, lifecycle_reason = state.lifecycle_reason, updated_at = state.updated_at
  FROM plan_po_state AS state WHERE po.id = state.id;
UPDATE procurement.purchase_request AS request
   SET commercial_plan_received_date = plan.received, version = request.version + 1, updated_at = now()
  FROM plan_request AS plan WHERE request.id = plan.id;""")
    parts.append(f"""INSERT INTO audit.audit_log (id, aggregate_type, aggregate_id, entity_type, entity_id, operation, field_name, new_value, actor_id, occurred_at, reason, correlation_id)
SELECT gen_random_uuid(), 'PURCHASE_REQUEST', id, 'PURCHASE_REQUEST', id, 'HISTORICAL_IMPORT', 'commercial_plan_received_date',
       jsonb_build_object('sha256', {sql_text(digest)}, 'date', received), {sql_text(ACTOR)}, now(),
       'Recebimento do plano comercial da planilha MRP.', {sql_text(batch_id)}::uuid
  FROM plan_request;
DO $reconcile$
DECLARE found bigint;
BEGIN
  SELECT count(*) INTO found FROM procurement.purchase_request WHERE commercial_plan_received_date IS NOT NULL;
  IF found <> {requests} THEN RAISE EXCEPTION 'Expected {requests} SCs with date, found %', found; END IF;
  SELECT count(*) INTO found FROM procurement.purchase_order_item WHERE commercial_plan_received_date IS NOT NULL;
  IF found <> {expected_items} THEN RAISE EXCEPTION 'Expected {expected_items} lines with date, found %', found; END IF;
  SELECT count(*) INTO found FROM plan_po_state AS state JOIN procurement.purchase_order AS po ON po.id = state.id
   WHERE ROW(po.lifecycle_status, po.lifecycle_changed_at, po.lifecycle_changed_by, po.lifecycle_reason, po.updated_at)
         IS DISTINCT FROM ROW(state.lifecycle_status, state.lifecycle_changed_at, state.lifecycle_changed_by, state.lifecycle_reason, state.updated_at);
  IF found <> 0 THEN RAISE EXCEPTION '% POs did not return to their state', found; END IF;
END $reconcile$;""")
    parts.append("COMMIT;" if commit else "ROLLBACK;")
    return "\n".join(parts) + "\n"


def read_workbook(path: Path) -> tuple[str, list[tuple[int, tuple]]]:
    from openpyxl import load_workbook

    digest = sha256(path.read_bytes()).hexdigest()
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        rows = workbook[SHEET].iter_rows(min_row=HEADER_ROW, max_col=SC_COLUMN + 1, values_only=True)
        check_headers(next(rows))
        return digest, [(index, row) for index, row in enumerate(rows, HEADER_ROW + 1)
                        if any(value not in (None, "") for value in row)]
    finally:
        workbook.close()


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--workbook", type=Path, required=True)
    parser.add_argument("--expected-sha256", required=True, help="SHA-256 of the workbook approved for this load")
    parser.add_argument("--expected-items", type=int, required=True, help="Lines that must end with a date; any other total stops the load")
    parser.add_argument("--sql", type=Path, required=True)
    parser.add_argument("--report", type=Path, required=True)
    parser.add_argument("--commit", action="store_true", help="End the SQL with COMMIT; the default is ROLLBACK, a rehearsal")
    args = parser.parse_args()
    digest, rows = read_workbook(args.workbook)
    if digest != args.expected_sha256:
        raise SystemExit(f"Workbook SHA-256 {digest} differs from the approved {args.expected_sha256}")
    plan = build_plan(rows)
    args.sql.write_text(render_sql(plan, digest, args.expected_items, args.commit), encoding="utf-8", newline="\n")
    report = summary(plan, digest, args.expected_items)
    args.report.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8", newline="\n")
    print(json.dumps({key: value for key, value in report.items() if key != "plan"}, ensure_ascii=False))


if __name__ == "__main__":
    main()
